"""Spreadsheet import. Append only.

The one rule this module exists to enforce: an import can add leads and it can
refresh their custom fields, but it can never delete a lead, never clear send
history, and never remove a suppression. There is no "replace list" path.
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field

from . import db

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s.]+\.[^@\s]+$")

EMAIL_HINTS = ("email", "e-mail", "mail", "email address", "contact email")

# What to do with an address already in the database.
ON_EXISTING = ("skip", "fill_blanks", "update")


@dataclass
class ParsedSheet:
    columns: list[str]
    rows: list[dict]
    filename: str


@dataclass
class ImportStats:
    new: int = 0
    existing: int = 0
    updated: int = 0
    suppressed: int = 0
    invalid: int = 0
    duplicate_in_file: int = 0
    invalid_samples: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "new": self.new,
            "existing": self.existing,
            "updated": self.updated,
            "suppressed": self.suppressed,
            "invalid": self.invalid,
            "duplicate_in_file": self.duplicate_in_file,
        }


def parse(data: bytes, filename: str) -> ParsedSheet:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(data, filename)
    return _parse_csv(data, filename)


def _parse_csv(data: bytes, filename: str) -> ParsedSheet:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    columns = [c.strip() for c in (reader.fieldnames or [])]
    rows = [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in reader]
    return ParsedSheet(columns=columns, rows=rows, filename=filename)


def _parse_xlsx(data: bytes, filename: str) -> ParsedSheet:
    from openpyxl import load_workbook  # imported lazily: CSV-only users skip the dependency

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None) or ()
    columns = [str(c).strip() if c is not None else "" for c in header]
    rows = []
    for raw in rows_iter:
        if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        row = {}
        for index, column in enumerate(columns):
            if not column:
                continue
            value = raw[index] if index < len(raw) else None
            row[column] = "" if value is None else str(value).strip()
        rows.append(row)
    workbook.close()
    return ParsedSheet(columns=columns, rows=rows, filename=filename)


def guess_email_column(columns: list[str]) -> str | None:
    lowered = {c.lower().strip(): c for c in columns if c}
    for hint in EMAIL_HINTS:
        if hint in lowered:
            return lowered[hint]
    for low, original in lowered.items():
        if "mail" in low:
            return original
    return None


def valid_email(value: str) -> bool:
    return bool(EMAIL_RE.match(value.strip().lower()))


def preview(sheet: ParsedSheet, email_column: str, limit: int = 10) -> tuple[ImportStats, list[dict]]:
    stats, prepared = _analyse(sheet, email_column)
    return stats, prepared[:limit]


def _analyse(sheet: ParsedSheet, email_column: str) -> tuple[ImportStats, list[dict]]:
    conn = db.connect()
    stats = ImportStats()
    seen: set[str] = set()
    prepared: list[dict] = []

    for row in sheet.rows:
        email = (row.get(email_column) or "").strip().lower()
        if not valid_email(email):
            stats.invalid += 1
            if len(stats.invalid_samples) < 5:
                stats.invalid_samples.append(email or "(blank)")
            continue
        if email in seen:
            stats.duplicate_in_file += 1
            continue
        seen.add(email)

        custom = {k: v for k, v in row.items() if k and k != email_column}
        exists = conn.execute("SELECT id FROM leads WHERE email=?", (email,)).fetchone()
        suppressed = db.is_suppressed(email)
        if suppressed:
            stats.suppressed += 1
        if exists:
            stats.existing += 1
        else:
            stats.new += 1
        prepared.append(
            {
                "email": email,
                "custom_fields": custom,
                "exists": bool(exists),
                "lead_id": exists["id"] if exists else None,
                "suppressed": suppressed,
            }
        )
    return stats, prepared


def ensure_list(name: str) -> int:
    conn = db.connect()
    row = conn.execute("SELECT id FROM lead_lists WHERE name=?", (name,)).fetchone()
    if row:
        return int(row["id"])
    cur = conn.execute(
        "INSERT INTO lead_lists (name, created_at) VALUES (?, ?)", (name, db.utcnow())
    )
    return int(cur.lastrowid)


def commit(sheet: ParsedSheet, email_column: str, list_name: str, on_existing: str = "skip") -> ImportStats:
    if on_existing not in ON_EXISTING:
        raise ValueError(f"on_existing must be one of {ON_EXISTING}")

    conn = db.connect()
    stats, prepared = _analyse(sheet, email_column)
    list_id = ensure_list(list_name)

    conn.execute("BEGIN")
    try:
        for item in prepared:
            if item["exists"]:
                lead_id = item["lead_id"]
                if on_existing != "skip":
                    current = json.loads(
                        conn.execute(
                            "SELECT custom_fields FROM leads WHERE id=?", (lead_id,)
                        ).fetchone()["custom_fields"]
                    )
                    merged = dict(current)
                    for key, value in item["custom_fields"].items():
                        if on_existing == "update" and value != "":
                            merged[key] = value
                        elif on_existing == "fill_blanks" and not current.get(key):
                            merged[key] = value
                    if merged != current:
                        conn.execute(
                            "UPDATE leads SET custom_fields=? WHERE id=?",
                            (json.dumps(merged), lead_id),
                        )
                        stats.updated += 1
            else:
                cur = conn.execute(
                    "INSERT INTO leads (email, custom_fields, source_file, created_at) "
                    "VALUES (?, ?, ?, ?)",
                    (
                        item["email"],
                        json.dumps(item["custom_fields"]),
                        sheet.filename,
                        db.utcnow(),
                    ),
                )
                lead_id = int(cur.lastrowid)
            # Membership is additive; a lead can belong to several lists.
            conn.execute(
                "INSERT OR IGNORE INTO lead_list_members (list_id, lead_id) VALUES (?, ?)",
                (list_id, lead_id),
            )
        conn.execute(
            "INSERT INTO imports (filename, list_id, stats, created_at) VALUES (?, ?, ?, ?)",
            (sheet.filename, list_id, json.dumps(stats.as_dict()), db.utcnow()),
        )
        conn.execute("COMMIT")
    except Exception:
        conn.execute("ROLLBACK")
        raise
    return stats


def custom_field_names(list_id: int | None = None) -> list[str]:
    conn = db.connect()
    if list_id:
        rows = conn.execute(
            "SELECT l.custom_fields FROM leads l "
            "JOIN lead_list_members m ON m.lead_id = l.id WHERE m.list_id=? LIMIT 500",
            (list_id,),
        ).fetchall()
    else:
        rows = conn.execute("SELECT custom_fields FROM leads LIMIT 500").fetchall()
    names: set[str] = set()
    for row in rows:
        names.update(json.loads(row["custom_fields"]).keys())
    return sorted(names)
